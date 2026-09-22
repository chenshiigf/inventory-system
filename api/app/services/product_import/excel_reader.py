from dataclasses import dataclass
from io import BytesIO
from typing import Final
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .schemas import IMPORT_COLUMNS, IMPORT_HEADERS


MAX_PREVIEW_ROWS: Final[int] = 20
IMPORT_SHEET_NAME: Final[str] = "商品导入"


class ProductImportWorkbookError(ValueError):
    """An uploaded workbook does not match the controlled import format."""


@dataclass(frozen=True)
class EmbeddedImage:
    data: bytes


@dataclass(frozen=True)
class RawImportRow:
    excel_row: int
    values: dict[str, object]
    images: tuple[EmbeddedImage, ...]


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _display_header(value: object) -> str:
    if _is_blank(value):
        return "空白"
    return str(value)


def _image_row(image: object) -> int | None:
    anchor = getattr(image, "anchor", None)
    from_cell = getattr(anchor, "_from", None)
    zero_based_row = getattr(from_cell, "row", None)
    if isinstance(zero_based_row, int) and zero_based_row >= 0:
        return zero_based_row + 1
    return None


def _read_embedded_images(worksheet) -> dict[int, list[EmbeddedImage]]:
    images_by_row: dict[int, list[EmbeddedImage]] = {}
    for image in getattr(worksheet, "_images", []):
        row = _image_row(image)
        if row is None:
            continue
        try:
            data = image._data()
        except (OSError, ValueError, AttributeError) as error:
            raise ProductImportWorkbookError("无法读取 Excel 中的商品图片。") from error
        if not data:
            raise ProductImportWorkbookError("无法读取 Excel 中的商品图片。")
        images_by_row.setdefault(row, []).append(EmbeddedImage(data=data))
    return images_by_row


def _validate_headers(worksheet) -> None:
    for column_index, expected_header in enumerate(IMPORT_HEADERS, start=1):
        actual_header = worksheet.cell(row=1, column=column_index).value
        if _display_header(actual_header) != expected_header:
            column_letter = get_column_letter(column_index)
            raise ProductImportWorkbookError(
                f"{column_letter}1 应为 {expected_header}，实际为 {_display_header(actual_header)}"
            )

    for column_index in range(len(IMPORT_COLUMNS) + 1, worksheet.max_column + 1):
        actual_header = worksheet.cell(row=1, column=column_index).value
        if not _is_blank(actual_header):
            column_letter = get_column_letter(column_index)
            raise ProductImportWorkbookError(
                f"{column_letter}1 不应有额外表头，实际为 {_display_header(actual_header)}"
            )


def read_import_workbook(data: bytes) -> list[RawImportRow]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as error:
        raise ProductImportWorkbookError(
            "无法读取 Excel 文件，请确认文件是有效的 .xlsx 工作簿。"
        ) from error

    try:
        if IMPORT_SHEET_NAME not in workbook.sheetnames:
            raise ProductImportWorkbookError(f"未找到工作表：{IMPORT_SHEET_NAME}")

        worksheet = workbook[IMPORT_SHEET_NAME]
        _validate_headers(worksheet)
        images_by_row = _read_embedded_images(worksheet)
        last_row = max(worksheet.max_row, max(images_by_row, default=1))
        rows: list[RawImportRow] = []

        for excel_row in range(2, last_row + 1):
            values = {
                column.key: worksheet.cell(row=excel_row, column=index).value
                for index, column in enumerate(IMPORT_COLUMNS, start=1)
            }
            images = tuple(images_by_row.get(excel_row, []))
            if all(_is_blank(value) for value in values.values()) and not images:
                continue

            rows.append(
                RawImportRow(
                    excel_row=excel_row,
                    values=values,
                    images=images,
                )
            )
            if len(rows) >= MAX_PREVIEW_ROWS:
                break

        return rows
    finally:
        workbook.close()
