from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schemas import IMPORT_COLUMNS


TEMPLATE_SHEET_NAME = "商品导入"
TEMPLATE_HEADER_FILL = "FF4F00"


def build_product_import_template() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = TEMPLATE_SHEET_NAME
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(IMPORT_COLUMNS))}1"
    worksheet.sheet_view.showGridLines = False

    worksheet.append([column.header for column in IMPORT_COLUMNS])
    for column_index, column in enumerate(IMPORT_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_index)
        cell.fill = PatternFill(fill_type="solid", fgColor=TEMPLATE_HEADER_FILL)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    comments = {
        "仓库": "填写系统中已经存在的仓库名称，例如：主仓。",
        "一级分类": "填写系统中已经存在的一级分类名称。",
        "二级分类": "填写属于前面一级分类的二级分类名称。",
        "产品图片": "允许为空。请将一张图片直接插入当前商品行的 D 列附近。",
        "产品尺寸": "允许为空，例如：20cm 或 20×20×5cm。只 trim 前后空格。",
        "装箱数": "必填，填写正整数，例如：24。不要填写 24pcs。",
        "单位": "仅支持 pcs 或 set，大小写会自动转为小写。",
        "单价": "必填，填写 >= 0 的数字，例如：3.50，不要带货币符号。",
        "当前箱数": "必填，填写 >= 0 的整数。",
        "备注": "允许为空，填写纯文本。",
        "原系统编号": "允许为空，仅用于迁移追踪和预览，不会生成商品编号。",
    }
    for column_index, column in enumerate(IMPORT_COLUMNS, start=1):
        worksheet.cell(row=1, column=column_index).comment = Comment(
            comments[column.header],
            "商品库存",
        )

    column_widths = {
        "A": 16,
        "B": 16,
        "C": 16,
        "D": 18,
        "E": 20,
        "F": 12,
        "G": 10,
        "H": 12,
        "I": 14,
        "J": 24,
        "K": 16,
    }
    for column_letter, width in column_widths.items():
        worksheet.column_dimensions[column_letter].width = width
    worksheet.row_dimensions[1].height = 26

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
