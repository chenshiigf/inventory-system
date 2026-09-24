"""Build an in-memory customer-facing product quote workbook."""

from __future__ import annotations

import re
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageOps, UnidentifiedImageError

from app.models import Product


SHEET_TITLE = "报价单"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BODY_FONT = Font(name="SimSun", size=12, color="1F2937")
_HEADER_FONT = Font(name="SimSun", size=12, bold=True, color="FFFFFF")
_TITLE_FONT = Font(name="SimSun", size=16, bold=True, color="1F3F68")
_META_FONT = Font(name="SimSun", size=12, color="4B5563")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="315A85")
_BORDER_SIDE = Side(style="thin", color="D9E2F0")
_CELL_BORDER = Border(
    left=_BORDER_SIDE,
    right=_BORDER_SIDE,
    top=_BORDER_SIDE,
    bottom=_BORDER_SIDE,
)
_MAX_IMAGE_EDGE = 105
_INVALID_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|]')


def build_quote_filename(customer_name: str | None, quote_date: date) -> str:
    normalized_name = _INVALID_FILENAME_CHARS.sub(
        "_", (customer_name or "").strip()
    )
    normalized_name = re.sub(r"\s+", " ", normalized_name).strip()
    name_part = f"_{normalized_name}" if normalized_name else ""
    return f"报价单{name_part}_{quote_date.isoformat()}.xlsx"


def build_quote_workbook(
    products: Sequence[Product],
    *,
    customer_name: str | None,
    quote_date: date,
    image_root: Path,
) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_TITLE
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A5"

    worksheet.merge_cells("A1:H1")
    title_cell = worksheet["A1"]
    title_cell.value = "商品报价单"
    title_cell.font = _TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30

    worksheet.merge_cells("A2:D2")
    worksheet.merge_cells("F2:H2")
    customer_cell = worksheet["A2"]
    customer_cell.value = f"客户名称：{customer_name or ''}"
    customer_cell.font = _META_FONT
    customer_cell.alignment = Alignment(vertical="center")
    date_cell = worksheet["F2"]
    date_cell.value = f"报价日期：{quote_date.isoformat()}"
    date_cell.font = _META_FONT
    date_cell.alignment = Alignment(horizontal="right", vertical="center")
    worksheet.row_dimensions[2].height = 23

    headers = ["序号", "产品图片", "商品编号", "尺寸", "装箱数", "单位", "单价", "备注"]
    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=4, column=column_index, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[4].height = 25

    column_widths = {
        "A": 7,
        "B": 18,
        "C": 16,
        "D": 20,
        "E": 18,
        "F": 10,
        "G": 12,
        "H": 32,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    for sequence, product in enumerate(products, start=1):
        row_number = sequence + 4
        values = [
            sequence,
            None,
            product.product_code or None,
            product.size.strip() or None,
            _format_packing_quantities(product),
            product.unit or None,
            _format_price(product.price),
            product.remark.strip() if product.remark and product.remark.strip() else None,
        ]
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column_index, value=value)
            cell.font = _BODY_FONT
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(
                horizontal="center" if column_index in {1, 2, 3, 5, 6, 7} else "left",
                vertical="center",
                wrap_text=column_index in {4, 8},
            )
            if column_index == 7 and value is not None:
                cell.number_format = "0.00"
        worksheet.row_dimensions[row_number].height = 90
        _add_product_image(worksheet, row_number, product, image_root)

    if products:
        worksheet.auto_filter.ref = f"A4:H{len(products) + 4}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _format_packing_quantities(product: Product) -> str | None:
    quantities = [
        str(packaging.packing_qty) if packaging.packing_qty is not None else ""
        for packaging in sorted(product.packagings, key=lambda item: item.sort_order)
    ]
    if not quantities or not any(quantities):
        return None
    return " / ".join(quantities)


def _format_price(price: Decimal | None) -> Decimal | None:
    return price if price is not None else None


def _add_product_image(
    worksheet,
    row_number: int,
    product: Product,
    image_root: Path,
) -> None:
    excel_image = None
    for stored_path in (product.image_path, product.thumbnail_path):
        image_path = _resolve_image_path(image_root, stored_path)
        if image_path is None:
            continue
        try:
            image_bytes = _read_image_as_png(image_path)
            excel_image = ExcelImage(BytesIO(image_bytes))
            break
        except (
            Image.DecompressionBombError,
            OSError,
            UnidentifiedImageError,
            ValueError,
        ):
            continue

    if excel_image is None:
        return

    width = max(1, int(excel_image.width))
    height = max(1, int(excel_image.height))
    scale = min(_MAX_IMAGE_EDGE / width, _MAX_IMAGE_EDGE / height, 1)
    excel_image.width = max(1, round(width * scale))
    excel_image.height = max(1, round(height * scale))
    worksheet.add_image(excel_image, f"B{row_number}")


def _resolve_image_path(image_root: Path, image_path: str | None) -> Path | None:
    if not image_path:
        return None

    normalized = image_path.strip().replace("\\", "/")
    relative_path = PurePosixPath(normalized)
    if (
        not normalized
        or relative_path.is_absolute()
        or ":" in normalized
        or any(part == ".." for part in relative_path.parts)
    ):
        return None

    root = image_root.resolve()
    candidate = (root / Path(*relative_path.parts)).resolve()
    try:
        candidate.relative_to(root)
    except ValueError:
        return None
    return candidate


def _read_image_as_png(image_path: Path) -> bytes:
    with Image.open(image_path) as source:
        oriented = ImageOps.exif_transpose(source)
        oriented.load()
        has_transparency = "A" in oriented.getbands() or "transparency" in oriented.info
        normalized = oriented.convert("RGBA" if has_transparency else "RGB")
        output = BytesIO()
        normalized.save(output, format="PNG")
        return output.getvalue()
