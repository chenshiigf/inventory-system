from collections.abc import Generator
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ElementTree

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import (
    AnchorMarker,
    OneCellAnchor,
    TwoCellAnchor,
)
from openpyxl.utils import get_column_letter
from PIL import Image as PillowImage
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.main import create_app
from app.models import Category, Product, ProductPackaging, Warehouse
from app.services.product_import.schemas import IMPORT_HEADERS, normalize_header


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_TO_FIELD = {
    "商品组": "product_group",
    "仓库": "warehouse",
    "一级分类": "category_level_1",
    "二级分类": "category_level_2",
    "产品图片": "image",
    "产品尺寸": "size",
    "装箱数": "packing_qty",
    "单位": "unit",
    "单价": "price",
    "当前箱数": "carton_count",
    "结余箱数": "carton_count",
    "备注": "remark",
    "原系统编号": "source_code",
    "系统编号": "source_code",
    "义库": "secondary_stock",
}
NORMALIZED_HEADER_TO_FIELD = {
    normalize_header(header): field for header, field in HEADER_TO_FIELD.items()
}


@dataclass(frozen=True)
class ImageSpec:
    data: bytes
    anchor: object | None = None


@pytest.fixture
def import_context(
    tmp_path: Path,
) -> Generator[tuple[TestClient, sessionmaker, FastAPI, Path, Path], None, None]:
    database_path = tmp_path / "product-import-test.db"
    uploads_directory = tmp_path / "uploads"
    preview_directory = tmp_path / "previews"
    engine = create_engine(
        f"sqlite:///{database_path.as_posix()}",
        connect_args={"check_same_thread": False},
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)
    application = create_app(
        uploads_directory=uploads_directory,
        product_import_preview_directory=preview_directory,
    )

    def override_get_db():
        with TestingSessionLocal() as db:
            yield db

    application.dependency_overrides[get_db] = override_get_db
    with TestClient(application) as client:
        yield client, TestingSessionLocal, application, uploads_directory, preview_directory

    application.dependency_overrides.clear()
    engine.dispose()


def seed_reference_data(session_factory: sessionmaker) -> dict[str, int]:
    with session_factory.begin() as db:
        main_warehouse = Warehouse(name="主仓", sort_order=0)
        tiger_warehouse = Warehouse(name="虎跳仓", sort_order=1)
        kitchen = Category(
            name="厨房用品",
            code="01",
            sort_order=0,
            next_product_sequence=None,
        )
        christmas = Category(
            name="圣诞系列",
            code="02",
            sort_order=1,
            next_product_sequence=None,
        )
        db.add_all([main_warehouse, tiger_warehouse, kitchen, christmas])
        db.flush()
        plate = Category(
            name="盘子",
            parent_id=kitchen.id,
            code="01",
            next_product_sequence=7,
        )
        bowl = Category(
            name="碗",
            parent_id=kitchen.id,
            code="02",
            next_product_sequence=1,
        )
        db.add_all([plate, bowl])
        db.flush()
        return {
            "main_warehouse": main_warehouse.id,
            "tiger_warehouse": tiger_warehouse.id,
            "kitchen": kitchen.id,
            "christmas": christmas.id,
            "plate": plate.id,
            "bowl": bowl.id,
        }


def image_bytes(
    color: tuple[int, int, int],
    *,
    width: int = 64,
    height: int = 25,
) -> bytes:
    output = BytesIO()
    PillowImage.new("RGB", (width, height), color).save(output, format="PNG")
    return output.getvalue()


def workbook_bytes(
    rows: list[dict[str, object] | None],
    *,
    images: dict[int, list[bytes | ImageSpec]] | None = None,
    sheet_name: str = "商品导入",
    headers: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    # Most regression fixtures emulate legacy workbooks that still contain
    # the optional source-code column. The downloaded template itself is
    # asserted separately and now contains only the 11 core headers.
    actual_headers = headers or [*IMPORT_HEADERS, "原系统编号"]
    worksheet.append(actual_headers)
    image_column_index = next(
        index
        for index, header in enumerate(actual_headers, start=1)
        if normalize_header(header) == "产品图片"
    )
    image_column_letter = get_column_letter(image_column_index)

    for excel_row, row in enumerate(rows, start=2):
        if row is not None:
            for column_index, header in enumerate(actual_headers, start=1):
                field = NORMALIZED_HEADER_TO_FIELD.get(normalize_header(header))
                value = row.get(field, row.get(header)) if field else row.get(header)
                worksheet.cell(row=excel_row, column=column_index).value = value
        for image_value in (images or {}).get(excel_row, []):
            spec = image_value if isinstance(image_value, ImageSpec) else ImageSpec(image_value)
            image = ExcelImage(BytesIO(spec.data))
            if spec.anchor is None:
                worksheet.add_image(image, f"{image_column_letter}{excel_row}")
            else:
                image.anchor = spec.anchor
                worksheet.add_image(image)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def with_formula_cache(
    data: bytes,
    *,
    cell_ref: str,
    formula: str,
    cached_value: object | None,
) -> bytes:
    source_files: dict[str, bytes] = {}
    with ZipFile(BytesIO(data)) as source:
        for name in source.namelist():
            source_files[name] = source.read(name)

    worksheet_name = next(
        name
        for name in source_files
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
    )
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    root = ElementTree.fromstring(source_files[worksheet_name])
    cell = next(
        node
        for node in root.iter(f"{{{namespace}}}c")
        if node.attrib.get("r") == cell_ref
    )
    cell.attrib.pop("t", None)
    for child in list(cell):
        if child.tag.rsplit("}", 1)[-1] in {"f", "v"}:
            cell.remove(child)
    formula_node = ElementTree.SubElement(cell, f"{{{namespace}}}f")
    formula_node.text = formula.removeprefix("=")
    if cached_value is not None:
        value_node = ElementTree.SubElement(cell, f"{{{namespace}}}v")
        value_node.text = str(cached_value)
    source_files[worksheet_name] = ElementTree.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
    )

    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name, content in source_files.items():
            target.writestr(name, content)
    return output.getvalue()


def valid_row(
    *,
    product_group: object = "",
    warehouse: object = "主仓",
    category_level_1: object = "厨房用品",
    category_level_2: object = "盘子",
    size: object = "20cm",
    packing_qty: object = 24,
    unit: object = "pcs",
    price: object = 3.50,
    carton_count: object = 10,
    remark: object = "蓝边款",
    source_code: object = "10235",
    secondary_stock: object | None = None,
) -> dict[str, object]:
    return {
        "product_group": product_group,
        "warehouse": warehouse,
        "category_level_1": category_level_1,
        "category_level_2": category_level_2,
        "image": None,
        "size": size,
        "packing_qty": packing_qty,
        "unit": unit,
        "price": price,
        "carton_count": carton_count,
        "remark": remark,
        "source_code": source_code,
        "secondary_stock": secondary_stock,
    }


def preview(
    client: TestClient,
    data: bytes,
    *,
    file_name: str = "厨房用品-导入.xlsx",
):
    return client.post(
        "/api/product-import/preview",
        files={"file": (file_name, data, XLSX_CONTENT_TYPE)},
    )


def one_product(response) -> dict[str, Any]:
    assert response.status_code == 200, response.text
    products = response.json()["products"]
    assert len(products) == 1
    return products[0]


def test_template_download_is_a_product_level_xlsx(import_context) -> None:
    client, _session_factory, _application, _uploads, _previews = import_context

    response = client.get("/api/product-import/template")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_CONTENT_TYPE)
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), read_only=False)
    worksheet = workbook["商品导入"]
    assert worksheet.max_row == 1
    assert [cell.value for cell in worksheet[1]] == list(IMPORT_HEADERS)
    assert all(
        worksheet.cell(row=1, column=index).comment
        for index in range(1, len(IMPORT_HEADERS) + 1)
    )
    assert worksheet[1][0].fill.fgColor.rgb.endswith("5B6B82")
    workbook.close()


def test_valid_preview_returns_one_product_and_packaging(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    data = workbook_bytes(
        [valid_row()],
        images={2: [image_bytes((90, 120, 160))]},
    )

    response = preview(client, data)

    product = one_product(response)
    payload = response.json()
    assert payload["file_name"] == "厨房用品-导入.xlsx"
    assert payload["source_row_count"] == 1
    assert payload["product_count"] == 1
    assert payload["valid_count"] == 1
    assert product["excel_rows"] == [2]
    assert product["warehouse"] == "主仓"
    assert product["category_level_1"] == "厨房用品"
    assert product["category_level_2"] == "盘子"
    assert product["packagings"] == [
        {
            "excel_row": 2,
            "packing_qty": 24,
            "carton_count": 10,
            "source_code": "10235",
        }
    ]
    assert product["total_carton_count"] == 10
    assert product["price"] == "3.50"
    assert product["status"] == "valid"
    assert product["shared_image"] is False
    assert product["image_preview_url"].startswith("/import-previews/")
    image_response = client.get(product["image_preview_url"])
    assert image_response.status_code == 200
    assert image_response.headers["content-type"] == "image/webp"


def test_headers_can_be_reordered_and_unknown_columns_are_ignored(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    headers = [
        "历史备注",
        "系统编号",
        "产品尺寸",
        "产品图片",
        "当前箱数",
        "商品组",
        "仓库",
        "一级分类",
        "二级分类",
        "装箱数",
        "单位",
        "单价",
        "备注",
    ]

    response = preview(
        client,
        workbook_bytes(
            [valid_row(source_code="SYS-01")],
            headers=headers,
            images={2: [image_bytes((40, 80, 120))]},
        ),
    )

    product = one_product(response)
    assert product["source_codes"] == ["SYS-01"]
    assert product["packagings"][0]["packing_qty"] == 24
    assert product["packagings"][0]["carton_count"] == 10
    assert product["status"] == "valid"


def test_stock_header_aliases_are_supported_and_both_are_rejected(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    alias_headers = [
        "商品组",
        "仓库",
        "一级分类",
        "二级分类",
        "产品图片",
        "产品尺寸",
        "装箱数",
        "单位",
        "单价",
        "结余箱数",
        "备注",
        "系统编号",
    ]
    alias_response = preview(client, workbook_bytes([valid_row()], headers=alias_headers))
    assert one_product(alias_response)["packagings"][0]["carton_count"] == 10

    missing_headers = [header for header in alias_headers if header != "结余箱数"]
    missing_response = preview(client, workbook_bytes([valid_row()], headers=missing_headers))
    assert missing_response.status_code == 400
    assert missing_response.json()["detail"] == "未找到可识别的商品数据工作表。"

    both_response = preview(
        client,
        workbook_bytes([valid_row()], headers=[*alias_headers, "当前箱数"]),
    )
    assert both_response.status_code == 400
    assert both_response.json()["detail"] == "同时发现‘结余箱数’和‘当前箱数’，请保留其中一个。"


@pytest.mark.parametrize(
    "header_variant",
    [
        "结余\n箱数",
        "结余 箱数",
        "结余　箱数",
        "备 注",
        "产品\n图片",
        "一 级 分 类",
    ],
)
def test_import_headers_are_normalized_across_all_aliases(
    import_context,
    header_variant: str,
) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    headers = list(IMPORT_HEADERS)
    normalized_variant = normalize_header(header_variant)
    target_header = (
        "当前箱数" if normalized_variant == "结余箱数" else normalized_variant
    )
    target_index = next(
        index
        for index, header in enumerate(headers)
        if normalize_header(header) == target_header
    )
    headers[target_index] = header_variant

    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row()],
                headers=headers,
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )

    assert product["status"] == "valid"
    assert product["packagings"][0]["carton_count"] == 10


def test_remark_internal_spaces_and_newlines_are_preserved(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    remark = "蓝边  大号\n第二行"

    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row(remark=f"  {remark}  ")],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )

    assert product["remark"] == remark


def test_non_default_sheet_name_is_selected_by_headers(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)

    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row()],
                sheet_name="库存明细",
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )

    assert product["status"] == "valid"


def test_multiple_recognizable_sheets_are_rejected(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    data = workbook_bytes(
        [valid_row()],
        sheet_name="库存明细",
        images={2: [image_bytes((90, 120, 160))]},
    )
    workbook = load_workbook(BytesIO(data))
    second_sheet = workbook.create_sheet("另一份库存")
    second_sheet.append(list(IMPORT_HEADERS))
    second_sheet.append([None] * len(IMPORT_HEADERS))
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    response = preview(client, output.getvalue())

    assert response.status_code == 400
    assert response.json()["detail"] == "发现多个可导入工作表：库存明细、另一份库存，请只保留一个。"


def test_optional_product_fields_and_single_packaging_quantity_can_be_blank(
    import_context,
) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row(unit=None, price=None, packing_qty=None)],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )

    assert product["status"] == "valid"
    assert product["unit"] is None
    assert product["price"] is None
    assert product["packagings"][0]["packing_qty"] is None


def test_multi_packaging_requires_nonblank_unique_packing_quantities(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [
                    valid_row(product_group="A001", packing_qty=None),
                    valid_row(product_group="A001", packing_qty=48),
                ],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )

    assert product["status"] == "error"
    assert "Excel第 2 行：装箱数不能为空" in product["messages"]


def test_blank_groups_remain_independent_products(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)

    response = preview(
        client,
        workbook_bytes(
            [
                valid_row(source_code="A"),
                valid_row(source_code="B", packing_qty=48),
            ]
        ),
    )

    payload = response.json()
    assert payload["source_row_count"] == 2
    assert payload["product_count"] == 2
    assert [product["excel_rows"] for product in payload["products"]] == [[2], [3]]


def test_same_product_group_merges_packagings_and_inherits_image(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [
        valid_row(product_group="A001", source_code="A-1", carton_count=10),
        {
            "product_group": "A001",
            "packing_qty": 48,
            "carton_count": 12,
            "source_code": "A-2",
        },
    ]

    response = preview(
        client,
        workbook_bytes(rows, images={2: [image_bytes((20, 80, 140))]}),
    )

    product = one_product(response)
    assert product["product_group"] == "A001"
    assert product["excel_rows"] == [2, 3]
    assert [item["packing_qty"] for item in product["packagings"]] == [24, 48]
    assert product["total_carton_count"] == 22
    assert product["source_codes"] == ["A-1", "A-2"]
    assert product["image_preview_url"]
    assert "无商品图片" not in product["messages"]
    assert product["warehouse"] == "主仓"
    assert product["category_level_1"] == "厨房用品"
    assert product["category_level_2"] == "盘子"
    assert product["size"] == "20cm"
    assert product["unit"] == "pcs"
    assert product["price"] == "3.50"
    assert product["remark"] == "蓝边款"


def test_blank_group_does_not_inherit_previous_product_fields(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    second = valid_row(
        warehouse="",
        category_level_1="",
        category_level_2="",
        source_code="B",
    )

    payload = preview(
        client,
        workbook_bytes([valid_row(source_code="A"), second]),
    ).json()

    assert payload["product_count"] == 2
    second_product = payload["products"][1]
    assert second_product["status"] == "error"
    assert "仓库不能为空" in second_product["messages"]
    assert "一级分类不能为空" in second_product["messages"]
    assert "二级分类不能为空" in second_product["messages"]


@pytest.mark.parametrize(
    ("field", "value", "expected"),
    [
        ("warehouse", "虎跳仓", "商品组 A001 的仓库不一致：主仓 / 虎跳仓"),
        ("category_level_1", "圣诞系列", "商品组 A001 的一级分类不一致：厨房用品 / 圣诞系列"),
        ("category_level_2", "碗", "商品组 A001 的二级分类不一致：盘子 / 碗"),
        ("size", "30cm", "商品组 A001 的产品尺寸不一致：20cm / 30cm"),
        ("unit", "set", "商品组 A001 的单位不一致：pcs / set"),
        ("price", 4.50, "商品组 A001 的单价不一致：3.50 / 4.50"),
        ("remark", "另一备注", "商品组 A001 的备注不一致"),
    ],
)
def test_same_group_field_conflicts_are_errors(
    import_context,
    field: str,
    value: object,
    expected: str,
) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    second = valid_row(product_group="A001", packing_qty=48)
    second[field] = value

    product = one_product(
        preview(client, workbook_bytes([valid_row(product_group="A001"), second]))
    )
    assert product["status"] == "error"
    assert expected in product["messages"]


@pytest.mark.parametrize(
    ("row_kwargs", "expected"),
    [
        ({"warehouse": "不存在仓"}, "仓库不存在：不存在仓"),
        ({"category_level_1": "不存在分类"}, "一级分类不存在：不存在分类"),
        ({"category_level_2": "不存在小类"}, "二级分类不存在：不存在小类"),
        (
            {"category_level_1": "圣诞系列", "category_level_2": "盘子"},
            "二级分类‘盘子’不属于一级分类‘圣诞系列’",
        ),
    ],
)
def test_warehouse_and_category_references_are_validated(import_context, row_kwargs, expected) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(preview(client, workbook_bytes([valid_row(**row_kwargs)])))
    assert product["status"] == "error"
    assert any(expected in message for message in product["messages"])


@pytest.mark.parametrize(
    ("row_kwargs", "expected"),
    [
        ({"unit": "box"}, "单位必须为 pcs 或 set"),
        ({"carton_count": -2}, "结余箱数不能小于 0"),
        ({"carton_count": 2.5}, "结余箱数必须是整数"),
        ({"packing_qty": 0}, "装箱数必须是正整数"),
        ({"packing_qty": "24pcs"}, "装箱数必须是整数"),
        ({"price": "¥3.50"}, "单价必须是 >= 0 且最多两位小数的合法数字"),
    ],
)
def test_product_values_are_validated(import_context, row_kwargs, expected) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(preview(client, workbook_bytes([valid_row(**row_kwargs)])))
    assert product["status"] == "error"
    assert any(expected in message for message in product["messages"])


@pytest.mark.parametrize(
    ("unit", "expected"),
    [("pcs", "pcs"), ("set", "set"), ("PCS", "pcs"), ("SET", "set")],
)
def test_allowed_units_are_normalized(import_context, unit: str, expected: str) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row(unit=unit)],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )
    assert product["unit"] == expected
    assert product["status"] == "valid"
    assert "无商品图片" not in product["messages"]


def test_zero_cartons_is_valid_and_missing_size_is_warning(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row(size="", carton_count=0)],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )
    assert product["packagings"][0]["carton_count"] == 0
    assert product["status"] == "warning"
    assert "产品尺寸为空" in product["messages"]
    assert "无商品图片" not in product["messages"]


def test_duplicate_packaging_quantity_is_an_error(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [
        valid_row(product_group="A001", source_code="A-1"),
        valid_row(product_group="A001", source_code="A-2"),
    ]
    product = one_product(preview(client, workbook_bytes(rows)))
    assert product["status"] == "error"
    assert "商品组 A001 存在重复装箱数：24" in product["messages"]
    assert len(product["packagings"]) == 1


def test_single_row_multiple_images_is_an_error(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row()],
                images={
                    2: [
                        image_bytes((10, 20, 30)),
                        image_bytes((30, 20, 10)),
                    ]
                },
            ),
        )
    )
    assert product["status"] == "error"
    assert any(
        "Excel第 2 行检测到多张商品图片" in message
        for message in product["messages"]
    )


def test_same_group_multiple_different_images_is_an_error(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [valid_row(product_group="A001"), valid_row(product_group="A001", packing_qty=48)]
    product = one_product(
        preview(
            client,
            workbook_bytes(
                rows,
                images={
                    2: [image_bytes((10, 20, 30))],
                    3: [image_bytes((30, 20, 10))],
                },
            ),
        )
    )
    assert product["status"] == "error"
    assert any(
        "商品组 A001 检测到多张商品图片" in message
        for message in product["messages"]
    )


def test_two_cell_anchor_shares_one_image_across_independent_products(import_context) -> None:
    client, session_factory, _application, _uploads, preview_directory = import_context
    seed_reference_data(session_factory)
    anchor = TwoCellAnchor(
        _from=AnchorMarker(col=4, row=1),
        to=AnchorMarker(col=4, row=2),
    )
    data = workbook_bytes(
        [valid_row(source_code="A"), valid_row(source_code="B", packing_qty=48)],
        images={2: [ImageSpec(image_bytes((70, 100, 150), height=55), anchor=anchor)]},
    )

    response = preview(client, data)

    payload = response.json()
    assert payload["product_count"] == 2
    products = payload["products"]
    assert [product["excel_rows"] for product in products] == [[2], [3]]
    assert all(product["image_preview_url"] for product in products)
    assert products[0]["image_preview_url"] == products[1]["image_preview_url"]
    assert products[0]["shared_image"] is True
    assert products[1]["shared_image"] is True
    assert all("无商品图片" not in product["messages"] for product in products)
    assert len(list(preview_directory.rglob("*.webp"))) == 1


def test_one_cell_anchor_requires_meaningful_coverage_for_next_row(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    anchor = OneCellAnchor(
        _from=AnchorMarker(col=4, row=1),
        ext=None,
    )
    # openpyxl serializes the image dimensions into ext when an image is added;
    # this image is deliberately just over one default Excel row.
    data = workbook_bytes(
        [valid_row(source_code="A"), valid_row(source_code="B", packing_qty=48)],
        images={2: [ImageSpec(image_bytes((70, 100, 150), height=25), anchor=anchor)]},
    )

    payload = preview(client, data).json()

    assert payload["products"][0]["image_preview_url"]
    assert payload["products"][1]["image_preview_url"] is None
    assert "无商品图片" in payload["products"][1]["messages"]


def test_one_cell_anchor_with_ext_covers_two_rows(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    anchor = OneCellAnchor(
        _from=AnchorMarker(col=4, row=1),
        ext=None,
    )
    data = workbook_bytes(
        [valid_row(source_code="A"), valid_row(source_code="B", packing_qty=48)],
        images={2: [ImageSpec(image_bytes((70, 100, 150), height=40), anchor=anchor)]},
    )

    payload = preview(client, data).json()

    products = payload["products"]
    assert payload["product_count"] == 2
    assert products[0]["image_preview_url"] == products[1]["image_preview_url"]
    assert products[0]["shared_image"] is True
    assert products[1]["shared_image"] is True
    assert all("无商品图片" not in product["messages"] for product in products)


def test_moved_image_column_is_detected_by_header(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    headers = [
        "商品组",
        "仓库",
        "一级分类",
        "二级分类",
        "产品尺寸",
        "装箱数",
        "单位",
        "单价",
        "当前箱数",
        "备注",
        "原系统编号",
        "产品图片",
    ]
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row()],
                headers=headers,
                images={2: [image_bytes((50, 120, 80))]},
            ),
        )
    )
    assert product["image_preview_url"]
    assert "无商品图片" not in product["messages"]


def test_secondary_stock_is_added_to_remark_and_bad_value_warns(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    headers = [*IMPORT_HEADERS, "义库"]
    rows = [
        valid_row(product_group="A001", secondary_stock=2),
        valid_row(product_group="A001", packing_qty=48, secondary_stock="无法识别"),
    ]
    product = one_product(
        preview(
            client,
            workbook_bytes(
                rows,
                headers=headers,
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )
    assert product["remark"].endswith("义库：2箱")
    assert product["status"] == "warning"
    assert "义库值无法识别，未合并到备注。" in product["messages"]


def test_formula_with_cached_result_is_read_as_number(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    data = workbook_bytes(
        [valid_row(carton_count="=G2+4")],
        images={2: [image_bytes((90, 120, 160))]},
    )
    data = with_formula_cache(
        data,
        cell_ref="J2",
        formula="=G2+4",
        cached_value=14,
    )

    product = one_product(preview(client, data))
    assert product["packagings"][0]["carton_count"] == 14
    assert product["status"] == "valid"
    assert all("没有可读取的计算结果" not in message for message in product["messages"])


def test_formula_without_cached_result_is_an_error(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    data = workbook_bytes(
        [valid_row(carton_count="=G2+4")],
        images={2: [image_bytes((90, 120, 160))]},
    )
    data = with_formula_cache(
        data,
        cell_ref="J2",
        formula="=G2+4",
        cached_value=None,
    )

    product = one_product(preview(client, data))
    assert product["status"] == "error"
    assert any(
        "结余箱数为公式，但没有可读取的计算结果。请使用 Microsoft Excel 打开文件、确认公式结果正常后保存，再重新上传。"
        in message
        for message in product["messages"]
    )


@pytest.mark.parametrize(
    ("cached_value", "expected_status", "expected_message"),
    [
        (0, "valid", None),
        (-2, "error", "结余箱数不能小于 0"),
        (2.5, "error", "结余箱数必须是整数"),
    ],
)
def test_formula_cached_values_follow_carton_count_rules(
    import_context,
    cached_value: object,
    expected_status: str,
    expected_message: str | None,
) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    data = workbook_bytes(
        [valid_row(carton_count="=G2+4")],
        images={2: [image_bytes((90, 120, 160))]},
    )
    data = with_formula_cache(
        data,
        cell_ref="J2",
        formula="=G2+4",
        cached_value=cached_value,
    )

    product = one_product(preview(client, data))
    assert product["status"] == expected_status
    if cached_value == 0:
        assert product["packagings"][0]["carton_count"] == 0
    if expected_message:
        assert any(expected_message in message for message in product["messages"])


def test_empty_size_is_warning_when_image_is_present(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    product = one_product(
        preview(
            client,
            workbook_bytes(
                [valid_row(size="", carton_count=0)],
                images={2: [image_bytes((90, 120, 160))]},
            ),
        )
    )
    assert product["status"] == "warning"
    assert product["packagings"][0]["carton_count"] == 0
    assert "产品尺寸为空" in product["messages"]
    assert "无商品图片" not in product["messages"]


def test_preview_skips_empty_rows_and_reads_all_source_rows(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows: list[dict[str, object] | None] = [valid_row(source_code="first")]
    rows.append({})
    rows.extend(valid_row(source_code=str(index)) for index in range(2, 22))

    response = preview(client, workbook_bytes(rows))

    payload = response.json()
    assert payload["source_row_count"] == 21
    assert [row["excel_rows"][0] for row in payload["products"]] == [2, *range(4, 24)]


def test_preview_reads_more_than_100_source_rows(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [valid_row(source_code=f"source-{index}") for index in range(130)]

    response = preview(client, workbook_bytes(rows))

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["source_row_count"] == 130
    assert len(payload["products"]) == 130


def test_preview_allows_300_source_rows(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [valid_row(source_code=f"source-{index}") for index in range(300)]

    response = preview(client, workbook_bytes(rows))

    assert response.status_code == 200, response.text
    assert response.json()["source_row_count"] == 300


def test_preview_rejects_more_than_300_source_rows(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)
    rows = [valid_row(source_code=f"source-{index}") for index in range(301)]

    response = preview(client, workbook_bytes(rows))

    assert response.status_code == 400
    assert response.json()["detail"] == "单个 Excel 最多支持 300 条数据行，请拆分后再导入。"


def test_product_without_any_image_is_an_error(import_context) -> None:
    client, session_factory, _application, _uploads, _previews = import_context
    seed_reference_data(session_factory)

    product = one_product(preview(client, workbook_bytes([valid_row()])))

    assert product["status"] == "error"
    assert "无商品图片" in product["messages"]


def test_duplicate_known_headers_are_rejected(import_context) -> None:
    client, _session_factory, _application, _uploads, _previews = import_context
    headers = [*IMPORT_HEADERS, "仓库"]
    response = preview(client, workbook_bytes([valid_row()], headers=headers))
    assert response.status_code == 400
    assert response.json()["detail"] == "发现重复字段表头：仓库、仓库，请保留其中一个。"


def test_preview_does_not_create_products_or_change_sequences_or_formal_uploads(
    import_context,
) -> None:
    client, session_factory, _application, uploads_directory, preview_directory = import_context
    ids = seed_reference_data(session_factory)
    with session_factory.begin() as db:
        existing_product = Product(
            category_id=ids["plate"],
            warehouse_id=ids["main_warehouse"],
            product_code="01-01-0007",
            size="existing",
            unit="pcs",
            price=Decimal("1.00"),
            remark="existing",
        )
        existing_product.packagings = [
            ProductPackaging(packing_qty=24, carton_count=6, sort_order=0)
        ]
        db.add(existing_product)
        db.flush()
        existing_packaging_id = existing_product.packagings[0].id
    data = workbook_bytes([valid_row()], images={2: [image_bytes((90, 120, 160))]})

    before_formal_files = list((uploads_directory / "products").rglob("*"))
    with session_factory() as db:
        before_product_count = db.scalar(select(func.count(Product.id)))
        before_packaging_count = db.scalar(select(func.count(ProductPackaging.id)))
        before_sequence = db.get(Category, ids["plate"]).next_product_sequence

    response = preview(client, data)

    assert response.status_code == 200
    with session_factory() as db:
        assert db.scalar(select(func.count(Product.id))) == before_product_count == 1
        assert (
            db.scalar(select(func.count(ProductPackaging.id)))
            == before_packaging_count
            == 1
        )
        existing_packaging = db.get(ProductPackaging, existing_packaging_id)
        assert existing_packaging is not None
        assert existing_packaging.carton_count == 6
        plate = db.get(Category, ids["plate"])
        assert plate is not None
        assert plate.next_product_sequence == before_sequence == 7
        assert (
            db.scalar(select(Product.product_code).where(Product.product_code.is_not(None)))
            == "01-01-0007"
        )

    after_formal_files = list((uploads_directory / "products").rglob("*"))
    assert after_formal_files == before_formal_files
    assert len(list(preview_directory.rglob("*.webp"))) == 1


@pytest.mark.parametrize(
    ("file_name", "expected_status"),
    [("legacy.xls", 415), ("legacy.csv", 415)],
)
def test_non_xlsx_files_are_rejected(import_context, file_name: str, expected_status: int) -> None:
    client, _session_factory, _application, _uploads, _previews = import_context
    response = preview(client, b"not an xlsx", file_name=file_name)
    assert response.status_code == expected_status


def test_invalid_xlsx_bytes_are_rejected(import_context) -> None:
    client, _session_factory, _application, _uploads, _previews = import_context
    response = preview(client, b"not an xlsx")
    assert response.status_code == 400
    assert "无法读取 Excel" in response.json()["detail"]


def test_workbook_without_recognizable_sheet_is_rejected(import_context) -> None:
    client, _session_factory, _application, _uploads, _previews = import_context
    response = preview(
        client,
        workbook_bytes(
            [valid_row()],
            sheet_name="Sheet1",
            headers=["名称", "产品图片"],
        ),
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "未找到可识别的商品数据工作表。"
