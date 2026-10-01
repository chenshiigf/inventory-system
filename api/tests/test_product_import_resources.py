"""Small sparse XLSX fixtures exercise resource limits without large files."""

import json
from copy import deepcopy
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select

from app.models import Product, ProductImportBatch
from app.services.product_import import excel_reader as reader, preview_service
from app.services.product_import.session_store import SOURCE_FILE_NAME, compute_file_hash
from tests.test_product_import import (
    ImageSpec, image_bytes, import_context, preview, seed_reference_data, valid_row,
    workbook_bytes, with_formula_cache,
)
from tests.test_product_import_tasks import normal_workbook, start, wait_terminal


S = reader.SHEET_NAMESPACE
D = reader.DRAWING_NAMESPACE


def rewrite_xml(data, transform):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename.lower().endswith(".xml"):
                root = ET.fromstring(content)
                transform(entry.filename, root)
                content = ET.tostring(root, encoding="utf-8")
            target.writestr(entry.filename, content)
    return output.getvalue()


def sparse_workbook(*, dimension="huge", data_row=None, image_row=None):
    data = workbook_bytes([valid_row()], images={2: [image_bytes((1, 2, 3))]} if image_row else None)

    def transform(name, root):
        if name.startswith("xl/worksheets/"):
            dim = root.find(f"{S}dimension")
            if dimension == "missing":
                root.remove(dim)
            else:
                dim.set("ref", "A1:XFD1048576" if dimension == "huge" else "A1:A1")
            sheet_data = root.find(f"{S}sheetData")
            if data_row:
                row = sheet_data.findall(f"{S}row")[1]
                row.set("r", str(data_row))
                for cell in row:
                    cell.set("r", cell.get("r").rstrip("0123456789") + str(data_row))
            # One serialized cell, not 1,048,576 x 16,384 cells.
            row = ET.SubElement(sheet_data, f"{S}row", r="1048576")
            ET.SubElement(row, f"{S}c", r="XFD1048576", s="0", t="n")
        if name.startswith("xl/drawings/") and image_row:
            for marker in root.iter(f"{D}from"):
                marker.find(f"{D}row").text = str(image_row - 1)

    return rewrite_xml(data, transform)


def huge_merge_workbook():
    def transform(name, root):
        if name.startswith("xl/worksheets/"):
            merges = ET.SubElement(root, f"{S}mergeCells", count="1")
            ET.SubElement(merges, f"{S}mergeCell", ref="A1:XFD1048576")
    return rewrite_xml(normal_workbook(), transform)


def forbid_rectangular_scans(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Reader must not generate blank row/column rectangles")
    monkeypatch.setattr(Worksheet, "iter_rows", forbidden)
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", forbidden)
    monkeypatch.setattr(ReadOnlyWorksheet, "calculate_dimension", forbidden)


@pytest.mark.parametrize("include_images", [False, True])
@pytest.mark.parametrize("dimension", ["huge", "missing", "too_small"])
def test_far_empty_formatting_and_bad_dimensions_do_not_scan_rectangles(monkeypatch, dimension, include_images):
    data = sparse_workbook(dimension=dimension)
    forbid_rectangular_scans(monkeypatch)
    rows = reader.read_import_workbook(data, include_images=include_images)
    assert [row.excel_row for row in rows] == [2]
    assert rows[0].values["packing_qty"] == 24
    assert rows[0].values["carton_count"] == 10


@pytest.mark.parametrize("include_images", [False, True])
def test_far_real_data_and_image_keep_original_row_without_blank_padding(monkeypatch, include_images):
    data = sparse_workbook(data_row=750000, image_row=750000)
    forbid_rectangular_scans(monkeypatch)
    rows = reader.read_import_workbook(data, include_images=include_images)
    assert [row.excel_row for row in rows] == [750000]
    assert rows[0].values["size"] == "20cm"
    assert bool(rows[0].images) == include_images


def test_far_image_only_tail_is_not_truncated_or_shifted(monkeypatch):
    data = sparse_workbook(image_row=900000)
    forbid_rectangular_scans(monkeypatch)
    rows = reader.read_import_workbook(data)
    assert [row.excel_row for row in rows] == [2, 900000]
    assert rows[1].images[0].covered_rows == (900000,)
    assert all(value is None for value in rows[1].values.values())


@pytest.mark.parametrize("include_images", [False, True])
def test_business_header_at_xfd_uses_actual_cells_not_sheet_width(monkeypatch, include_images):
    def transform(name, root):
        if name.startswith("xl/worksheets/"):
            for row in root.find(f"{S}sheetData"):
                for cell in row:
                    if cell.get("r", "").startswith("G"):
                        cell.set("r", "XFD" + row.get("r"))
    data = rewrite_xml(workbook_bytes([valid_row()]), transform)
    forbid_rectangular_scans(monkeypatch)
    rows = reader.read_import_workbook(data, include_images=include_images)
    assert rows[0].values["packing_qty"] == 24


@pytest.mark.parametrize("include_images", [False, True])
@pytest.mark.parametrize("count", [300, 301])
def test_existing_nonempty_row_business_limit_is_unchanged(count, include_images):
    data = workbook_bytes([valid_row() for _ in range(count)])
    if count == 300:
        rows = reader.read_import_workbook(data, include_images=include_images)
        assert len(rows) == 300
        assert rows[-1].excel_row == 301
    else:
        with pytest.raises(reader.ProductImportWorkbookError, match=reader.TOO_MANY_SOURCE_ROWS_MESSAGE):
            reader.read_import_workbook(data, include_images=include_images)


def test_small_workbook_preserves_cached_formulas_and_blank_text():
    data = with_formula_cache(workbook_bytes([valid_row(remark="   ")]), cell_ref="J2", formula="4+6", cached_value=10)
    rows = reader.read_import_workbook(data, include_images=False)
    assert rows[0].excel_row == 2
    assert rows[0].values["remark"] == "   "
    assert rows[0].values["carton_count"] == 10
    assert rows[0].formulas == {"carton_count": "=4+6"}


@pytest.mark.parametrize("limit", [
    "MAX_ZIP_ENTRIES", "MAX_ZIP_DIRECTORY_BYTES", "MAX_EXPANDED_BYTES", "MAX_XML_BYTES", "MAX_XML_PART_BYTES",
    "MAX_XML_ELEMENTS", "MAX_SERIALIZED_ROWS", "MAX_WORKBOOK_CELLS", "MAX_WORKBOOK_IMAGES",
])
def test_resource_budgets_fail_before_workbook_or_image_loading(monkeypatch, limit):
    # Lower only the tested budget; the fixture itself stays a tiny real XLSX.
    data = normal_workbook()
    monkeypatch.setattr(reader, limit, 0)
    def forbidden(*args, **kwargs):
        pytest.fail("Resource rejection must precede openpyxl/Pillow loading")
    monkeypatch.setattr(reader, "load_workbook", forbidden)
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data)


@pytest.mark.parametrize("limit", ["MAX_ZIP_ENTRIES", "MAX_ZIP_DIRECTORY_BYTES"])
def test_zip_directory_budget_is_checked_before_allocating_entries(monkeypatch, limit):
    data = normal_workbook()
    monkeypatch.setattr(reader, limit, 0)
    monkeypatch.setattr(reader, "ZipFile", lambda *a, **kw: pytest.fail("Must check ZIP end record before allocating entries"))
    with pytest.raises(reader.ProductImportWorkbookError, match="压缩包目录异常过大"):
        reader.read_import_workbook(data)


def test_huge_merge_is_rejected_before_openpyxl_materializes_cells(monkeypatch):
    data = huge_merge_workbook()
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must reject giant merge first"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data, include_images=False)


def test_huge_hyperlink_range_is_rejected_before_materialization(monkeypatch):
    def transform(name, root):
        if name.startswith("xl/worksheets/"):
            links = ET.SubElement(root, f"{S}hyperlinks")
            ET.SubElement(links, f"{S}hyperlink", ref="A1:XFD1048576", location="A1")
    data = rewrite_xml(normal_workbook(), transform)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must reject hyperlink range first"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data)


def test_comment_range_cannot_materialize_a_rectangle(monkeypatch):
    output = BytesIO()
    with ZipFile(BytesIO(normal_workbook())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, source.read(entry))
        target.writestr("xl/comments/comment1.xml", (
            '<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<commentList><comment ref="A1:XFD1048576" authorId="0"/></commentList></comments>'
        ))
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must reject invalid comment range first"))
    with pytest.raises(reader.ProductImportWorkbookError, match="有效的 .xlsx"):
        reader.read_import_workbook(output.getvalue())


@pytest.mark.parametrize("kind", ["span", "coordinate", "associations"])
def test_abnormal_image_anchors_are_rejected_during_precheck(monkeypatch, kind):
    anchor = TwoCellAnchor(_from=AnchorMarker(col=4, row=1), to=AnchorMarker(col=4, row=100000))
    if kind == "coordinate":
        anchor.to.row = reader.EXCEL_MAX_ROW
    elif kind == "associations":
        anchor.to.row = 3
        monkeypatch.setattr(reader, "MAX_IMAGE_ROW_ASSOCIATIONS", 2)
    data = workbook_bytes([valid_row()], images={2: [ImageSpec(image_bytes((1, 2, 3)), anchor=anchor)]})
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must check anchors before loading images"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大|锚点超出"):
        reader.read_import_workbook(data, include_images=False)


def test_one_cell_huge_extent_is_rejected_before_images(monkeypatch):
    def transform(name, root):
        if name.startswith("xl/drawings/"):
            for node in root.iter(f"{D}ext"):
                node.set("cy", "999999999999999999")
    data = rewrite_xml(normal_workbook(), transform)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must check extent first"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data, include_images=False)


def test_invalid_xml_is_reported_as_workbook_error():
    output = BytesIO()
    with ZipFile(BytesIO(normal_workbook())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, b"<worksheet>" if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry))
    with pytest.raises(reader.ProductImportWorkbookError, match="有效的 .xlsx"):
        reader.read_import_workbook(output.getvalue())


@pytest.mark.parametrize("endpoint", ["legacy", "task"])
def test_resource_error_is_user_error_and_leaves_no_task_or_preview(import_context, monkeypatch, endpoint):
    client, factory, app, uploads, directory = import_context
    seed_reference_data(factory)
    with monkeypatch.context() as patch:
        patch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Must reject before image loading"))
        patch.setattr(preview_service, "_save_preview_image", lambda *a, **kw: pytest.fail("Must not process images"))
        result = (preview if endpoint == "legacy" else start)(client, huge_merge_workbook())
    assert result.status_code == 400
    assert "异常过大" in result.json()["detail"]
    assert not list(directory.iterdir())
    assert not list(uploads.rglob("*.webp"))
    assert not app.state.product_import_preview_tasks._tasks
    # Reservation is released and a normal subsequent upload still works.
    task_id = start(client, normal_workbook()).json()["task_id"]
    assert wait_terminal(client, task_id)["status"] == "succeeded"


def test_background_worker_rechecks_same_resource_limits(import_context, monkeypatch):
    client, factory, app, _, directory = import_context
    seed_reference_data(factory)
    registry = app.state.product_import_preview_tasks
    real_start = registry.start
    def replace_source_before_worker(**kwargs):
        (kwargs["directory"] / SOURCE_FILE_NAME).write_bytes(huge_merge_workbook())
        return real_start(**kwargs)
    monkeypatch.setattr(registry, "start", replace_source_before_worker)
    monkeypatch.setattr(preview_service, "_save_preview_image", lambda *a, **kw: pytest.fail("Invalid resources must not reach image processing"))
    response = start(client, normal_workbook())
    assert response.status_code == 202
    task = wait_terminal(client, response.json()["task_id"])
    assert task["status"] == "failed"
    assert "异常过大" in task["error"]
    assert not list(directory.iterdir())


def test_commit_reread_uses_same_limits_and_writes_no_products(import_context, monkeypatch):
    client, factory, _, uploads, directory = import_context
    seed_reference_data(factory)
    result = preview(client, normal_workbook())
    assert result.status_code == 200
    session_id = result.json()["preview_session_id"]
    session = directory / session_id
    # Emulate a pre-existing session containing a now-disallowed resource file.
    # Use a consistent hash so the test reaches the reader, not hash validation.
    data = huge_merge_workbook()
    (session / SOURCE_FILE_NAME).write_bytes(data)
    metadata_path = session / "preview.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata["file_hash"] = compute_file_hash(data)
    metadata_path.write_text(json.dumps(metadata), encoding="utf-8")
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Commit must reject before image loading"))
    response = client.post("/api/product-import/commit", json={"preview_session_id": session_id})
    assert response.status_code == 409
    assert "原始 Excel 无法重新校验" in response.json()["detail"]
    assert "异常过大" in response.json()["detail"]
    with factory() as db:
        assert db.scalar(select(func.count(Product.id))) == 0
        assert db.scalar(select(func.count(ProductImportBatch.id))) == 0
    assert not list(uploads.rglob("*.webp"))


def renamed_parts(data, replacements):
    """Rename exact ZIP parts and synchronize package/relationship targets."""
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename.lower().endswith((".xml", ".rels")):
                for old, new in replacements.items():
                    content = content.replace(old.encode(), new.encode())
                    content = content.replace(old.rsplit("/", 1)[-1].encode(), new.rsplit("/", 1)[-1].encode())
            target.writestr(replacements.get(entry.filename, entry.filename), content)
    return output.getvalue()


def shared_drawing_workbook(sheets, anchors=1, drawing_references=1):
    """Reuse one drawing/media using small XML copies; no large image fixture."""
    data = normal_workbook()
    package_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    document_ns = reader.DOCUMENT_RELATIONSHIP_NAMESPACE
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        workbook = ET.fromstring(source.read("xl/workbook.xml"))
        workbook_sheets = workbook.find(f"{S}sheets")
        workbook_rels = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        types = ET.fromstring(source.read("[Content_Types].xml"))
        sheet = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        empty_sheet = deepcopy(sheet)
        empty_sheet.find(f"{S}sheetData").clear()
        empty_sheet.find(f"{S}dimension").set("ref", "A1:A1")
        sheet_rels = ET.fromstring(source.read("xl/worksheets/_rels/sheet1.xml.rels"))
        for index in range(1, drawing_references):
            relation = deepcopy(sheet_rels[0])
            relation.set("Id", f"extraDrawing{index}")
            sheet_rels.append(relation)
        drawing = ET.fromstring(source.read("xl/drawings/drawing1.xml"))
        original_anchor = deepcopy(drawing[0])
        drawing.clear()
        for index in range(anchors):
            anchor = deepcopy(original_anchor)
            for properties in anchor.iter(f"{D}cNvPr"):
                properties.set("id", str(index + 1))
            drawing.append(anchor)
        for index in range(2, sheets + 1):
            relation_id = f"sharedSheet{index}"
            ET.SubElement(workbook_sheets, f"{S}sheet", {
                "name": f"共享图片{index}", "sheetId": str(index), f"{document_ns}id": relation_id,
            })
            ET.SubElement(workbook_rels, f"{package_ns}Relationship", {
                "Id": relation_id, "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
                "Target": f"/xl/worksheets/sheet{index}.xml",
            })
            ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
                "PartName": f"/xl/worksheets/sheet{index}.xml",
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
            })
            target.writestr(f"xl/worksheets/sheet{index}.xml", ET.tostring(empty_sheet))
            target.writestr(f"xl/worksheets/_rels/sheet{index}.xml.rels", ET.tostring(sheet_rels))
        rewritten = {
            "xl/workbook.xml": ET.tostring(workbook), "xl/_rels/workbook.xml.rels": ET.tostring(workbook_rels),
            "[Content_Types].xml": ET.tostring(types), "xl/drawings/drawing1.xml": ET.tostring(drawing),
            "xl/worksheets/_rels/sheet1.xml.rels": ET.tostring(sheet_rels),
        }
        for entry in source.infolist():
            target.writestr(entry.filename, rewritten.get(entry.filename, source.read(entry)))
    return output.getvalue()


@pytest.mark.parametrize("suffix", ["XML", "xMl", "XmL", "bin"])
def test_renamed_xml_huge_merge_is_rejected_before_materialization(monkeypatch, suffix):
    data = renamed_parts(huge_merge_workbook(), {"xl/worksheets/sheet1.xml": f"xl/worksheets/sheet1.{suffix}"})
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Renamed merge must be rejected before openpyxl"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data, include_images=False)


@pytest.mark.parametrize("suffix", ["XML", "xMl"])
@pytest.mark.parametrize("limit", ["MAX_XML_PART_BYTES", "MAX_XML_BYTES", "MAX_XML_ELEMENTS"])
def test_uppercase_xml_is_in_all_xml_budgets(monkeypatch, suffix, limit):
    data = renamed_parts(normal_workbook(), {"xl/worksheets/sheet1.xml": f"xl/worksheets/sheet1.{suffix}"})
    with ZipFile(BytesIO(data)) as archive:
        budget = (max(e.file_size for e in archive.infolist() if e.filename.endswith((".xml", ".rels")))
                  if limit == "MAX_XML_PART_BYTES" else
                  sum(e.file_size for e in archive.infolist() if e.filename.endswith((".xml", ".rels"))))
    if limit == "MAX_XML_PART_BYTES":
        # Make the renamed part larger than every other XML part without many nodes.
        def grow(name, root):
            if name.endswith(f".{suffix}"):
                root.set("padding", "x" * (budget + 1))
        data = rewrite_xml(data, grow)
    elif limit == "MAX_XML_ELEMENTS":
        with ZipFile(BytesIO(data)) as archive:
            budget = sum(sum(1 for _ in ET.fromstring(archive.read(e)).iter()) for e in archive.infolist() if e.filename.endswith((".xml", ".rels")))
    monkeypatch.setattr(reader, limit, budget)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("All XML budgets must include uppercase parts"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data)


@pytest.mark.parametrize("suffix", ["RELS", "ReLs"])
def test_relationship_case_is_scanned_for_xml_nodes(monkeypatch, suffix):
    data = renamed_parts(normal_workbook(), {"xl/worksheets/_rels/sheet1.xml.rels": f"xl/worksheets/_rels/sheet1.xml.{suffix}"})
    with ZipFile(BytesIO(data)) as archive:
        nodes = sum(sum(1 for _ in ET.fromstring(archive.read(e)).iter()) for e in archive.infolist() if e.filename.endswith((".xml", ".rels")))
    monkeypatch.setattr(reader, "MAX_XML_ELEMENTS", nodes)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Uppercase relationships must be scanned"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data, include_images=False)


@pytest.mark.parametrize("sheets,anchors,refs", [(2, 1, 1), (10, 100, 1), (1, 500, 2)])
def test_legal_shared_drawing_and_media_within_budget_load_normally(monkeypatch, sheets, anchors, refs):
    data = shared_drawing_workbook(sheets, anchors, refs)
    real_load = reader.load_workbook
    loaded_images = []
    def observe(*args, **kwargs):
        workbook = real_load(*args, **kwargs)
        if not workbook.read_only:
            loaded_images.append(sum(len(sheet._images) for sheet in workbook))
        return workbook
    monkeypatch.setattr(reader, "load_workbook", observe)
    rows = reader.read_import_workbook(data)
    assert loaded_images == [sheets * anchors * refs]
    assert [row.excel_row for row in rows] == [2]
    assert len(rows[0].images) == anchors * refs


@pytest.mark.parametrize("sheets,anchors,refs", [(20, 100, 1), (1, 501, 2), (1001, 1, 1)])
def test_shared_drawing_expansion_over_1000_fails_before_loading(monkeypatch, sheets, anchors, refs):
    data = shared_drawing_workbook(sheets, anchors, refs)
    assert len(data) < 1024 * 1024
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Expanded resource costs must be checked first"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data, include_images=False)


def test_same_media_repeated_anchors_counts_actual_bytes(monkeypatch):
    data = shared_drawing_workbook(2, 2)
    with ZipFile(BytesIO(data)) as archive:
        unique_bytes = sum(e.file_size for e in archive.infolist())
        repeated_xml = archive.getinfo("xl/drawings/drawing1.xml").file_size + archive.getinfo("xl/drawings/_rels/drawing1.xml.rels").file_size
    monkeypatch.setattr(reader, "MAX_EXPANDED_BYTES", unique_bytes + repeated_xml)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Repeated media/drawing reads must count"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data)


def test_shared_media_expanded_bytes_exact_boundary_is_allowed(monkeypatch):
    data = shared_drawing_workbook(2, 2)
    with ZipFile(BytesIO(data)) as archive:
        unique_bytes = sum(e.file_size for e in archive.infolist())
        repeated_xml = archive.getinfo("xl/drawings/drawing1.xml").file_size + archive.getinfo("xl/drawings/_rels/drawing1.xml.rels").file_size
        repeated_media = 3 * archive.getinfo("xl/media/image1.png").file_size
    monkeypatch.setattr(reader, "MAX_EXPANDED_BYTES", unique_bytes + repeated_xml + repeated_media)
    reader._check_workbook_resources(data, lambda: None)


def test_mixed_case_drawing_and_relationship_expansion_is_checked(monkeypatch):
    data = renamed_parts(shared_drawing_workbook(20, 100), {
        "xl/drawings/drawing1.xml": "xl/drawings/drawing1.XmL",
        "xl/drawings/_rels/drawing1.xml.rels": "xl/drawings/_rels/drawing1.XmL.ReLs",
    })
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Mixed-case references must be budgeted before loading"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(data)


@pytest.mark.parametrize("part", ["relationships", "drawing"])
def test_reference_costs_match_openpyxl_namespace_tolerant_elements(monkeypatch, part):
    data = shared_drawing_workbook(20, 100)
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry)
            if (part == "relationships" and entry.filename.endswith(".rels")) or (part == "drawing" and entry.filename == "xl/drawings/drawing1.xml"):
                root = ET.fromstring(content)
                for node in root.iter():
                    node.tag = node.tag.rsplit("}", 1)[-1]
                content = ET.tostring(root)
            target.writestr(entry.filename, content)
    monkeypatch.setattr(reader, "load_workbook", lambda *a, **kw: pytest.fail("Namespace variations must not disable reference counting"))
    with pytest.raises(reader.ProductImportWorkbookError, match="异常过大"):
        reader.read_import_workbook(output.getvalue())
