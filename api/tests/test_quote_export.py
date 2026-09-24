from collections.abc import Generator
from datetime import date
from io import BytesIO
from urllib.parse import unquote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import selectinload, sessionmaker

from app.database import Base, get_db
from app.main import app
from app.models import InventoryMovement, Product
from app.services.quote_export import build_quote_workbook


@pytest.fixture
def database_client(tmp_path) -> Generator[tuple[TestClient, sessionmaker], None, None]:
    database_path = tmp_path / "quote-export-test.db"
    engine = create_engine(
        f"sqlite:///{database_path.as_posix()}",
        connect_args={"check_same_thread": False},
    )
    testing_session_local = sessionmaker(
        bind=engine,
        autoflush=False,
        autocommit=False,
    )
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        with testing_session_local() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client, testing_session_local

    app.dependency_overrides.clear()
    engine.dispose()


def create_product(client: TestClient, **overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "image_path": None,
        "size": "18 × 18 cm",
        "packagings": [{"packing_qty": 240, "carton_count": 3}],
        "unit": "pcs",
        "price": "2.80",
        "remark": "蓝边方盘",
    }
    payload.update(overrides)
    response = client.post("/api/products", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def set_product_fields(session_local, product_id: int, **fields: object) -> None:
    with session_local.begin() as db:
        product = db.get(Product, product_id)
        assert product is not None
        for field_name, value in fields.items():
            setattr(product, field_name, value)


def read_workbook(response_content: bytes):
    workbook = load_workbook(BytesIO(response_content), data_only=False)
    return workbook, workbook["报价单"]


def test_batch_quote_export_returns_ordered_customer_workbook(database_client) -> None:
    client, session_local = database_client
    first = create_product(
        client,
        size="5*2.5*7cm",
        packagings=[{"packing_qty": 240, "carton_count": 3}],
        unit="pcs",
        price="1.50",
        remark="第一件备注",
    )
    second = create_product(
        client,
        size="",
        packagings=[
            {"packing_qty": 144, "carton_count": 1},
            {"packing_qty": 256, "carton_count": 2},
            {"packing_qty": 96, "carton_count": 4},
        ],
        unit=None,
        price=None,
        remark=None,
    )
    set_product_fields(
        session_local,
        first["id"],
        product_code="01-03-014",
    )
    set_product_fields(
        session_local,
        second["id"],
        product_code="01-03-013",
    )

    response = client.post(
        "/api/products/batch/export-quote",
        json={
            "product_ids": [second["id"], first["id"], second["id"]],
            "customer_name": "  客户 / A  ",
            "quote_date": "2026-09-24",
        },
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    content_disposition = response.headers["content-disposition"]
    encoded_filename = content_disposition.split("filename*=UTF-8''", 1)[1]
    assert unquote(encoded_filename).startswith("报价单_客户 _ A_")

    workbook, worksheet = read_workbook(response.content)
    assert worksheet["A1"].value == "商品报价单"
    assert worksheet["A2"].value == "客户名称：客户 / A"
    assert worksheet["F2"].value == "报价日期：2026-09-24"
    assert [worksheet.cell(4, column).value for column in range(1, 9)] == [
        "序号",
        "产品图片",
        "商品编号",
        "尺寸",
        "装箱数",
        "单位",
        "单价",
        "备注",
    ]
    assert [worksheet.cell(row, 3).value for row in (5, 6)] == [
        "01-03-013",
        "01-03-014",
    ]
    assert worksheet["D5"].value is None
    assert worksheet["E5"].value == "144 / 256 / 96"
    assert worksheet["F5"].value is None
    assert worksheet["G5"].value is None
    assert worksheet["H5"].value is None
    assert worksheet["D6"].value == "5*2.5*7cm"
    assert worksheet["E6"].value == "240"
    assert worksheet["F6"].value == "pcs"
    assert worksheet["G6"].value == 1.5
    assert worksheet["G6"].data_type == "n"
    assert worksheet["H6"].value == "第一件备注"
    assert worksheet.max_row == 6
    assert worksheet.freeze_panes == "A5"
    workbook.close()


def test_batch_quote_export_keeps_null_packing_qty_blank(database_client) -> None:
    client, _session_local = database_client
    product = create_product(
        client,
        packagings=[{"packing_qty": None, "carton_count": 3}],
    )

    response = client.post(
        "/api/products/batch/export-quote",
        json={"product_ids": [product["id"]]},
    )

    assert response.status_code == 200, response.text
    workbook, worksheet = read_workbook(response.content)
    assert worksheet["E5"].value is None
    workbook.close()


def test_batch_quote_export_allows_inactive_and_does_not_mutate_inventory(
    database_client,
) -> None:
    client, session_local = database_client
    product = create_product(client)
    set_product_fields(session_local, product["id"], is_active=False)
    before = client.get(f"/api/products/{product['id']}").json()

    response = client.post(
        "/api/products/batch/export-quote",
        json={"product_ids": [product["id"]]},
    )

    assert response.status_code == 200, response.text
    after = client.get(f"/api/products/{product['id']}").json()
    assert after["is_active"] is False
    assert after["packagings"] == before["packagings"]
    assert after["total_carton_count"] == before["total_carton_count"]
    with session_local() as db:
        assert db.scalar(select(func.count(InventoryMovement.id))) == 0


@pytest.mark.parametrize(
    ("payload", "expected_status"),
    [
        ({"product_ids": []}, 422),
        ({"product_ids": [1] * 101}, 422),
        ({"product_ids": [999999]}, 404),
    ],
)
def test_batch_quote_export_validates_all_ids_before_generating(
    database_client,
    payload: dict[str, object],
    expected_status: int,
) -> None:
    client, _session_local = database_client

    response = client.post("/api/products/batch/export-quote", json=payload)

    assert response.status_code == expected_status
    assert response.headers.get("content-type", "").startswith("application/json")


def test_batch_quote_export_embeds_webp_without_distorting_aspect_ratio(
    database_client,
    tmp_path,
) -> None:
    client, session_local = database_client
    product_data = create_product(client)
    image_root = tmp_path / "uploads"
    image_path = image_root / "products" / "main" / "quote-image.webp"
    image_path.parent.mkdir(parents=True)
    with Image.new("RGBA", (240, 120), (20, 100, 180, 180)) as image:
        image.save(image_path, format="WEBP")
    set_product_fields(
        session_local,
        product_data["id"],
        image_path="products/main/quote-image.webp",
    )

    with session_local() as db:
        product = db.scalar(
            select(Product)
            .options(selectinload(Product.packagings))
            .where(Product.id == product_data["id"])
        )
        assert product is not None
        workbook_stream = build_quote_workbook(
            [product],
            customer_name=None,
            quote_date=date(2026, 9, 24),
            image_root=image_root,
        )

    workbook, worksheet = read_workbook(workbook_stream.getvalue())
    assert len(worksheet._images) == 1
    embedded_image = worksheet._images[0]
    assert embedded_image.anchor.ext.width == 105 * 9525
    assert embedded_image.anchor.ext.height == 52 * 9525
    workbook.close()
